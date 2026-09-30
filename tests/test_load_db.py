import datetime as dt
import io
import json
import sys
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import load_db  # noqa: E402

POINT = {"type": "Point", "coordinates": [23.32, 42.69]}
AREA = {"type": "MultiPolygon", "coordinates": [[[[23.3, 42.6], [23.4, 42.6], [23.4, 42.7], [23.3, 42.6]]]]}


def collection(*features, **extra):
    return dict({"type": "FeatureCollection", "features": list(features)}, **extra)


class FeatureRowsTest(unittest.TestCase):
    def rows(self, doc):
        stats = load_db.LayerStats()
        rows = list(load_db.feature_rows(doc))
        for _, props, geom in rows:
            stats.add(props, geom)
        return rows, stats

    def test_rows_and_stats(self):
        rows, stats = self.rows(collection(
            {"type": "Feature", "id": 7, "properties": {"name": "a", "n": 1}, "geometry": AREA},
            {"type": "Feature", "properties": {"name": None, "n": 2.5, "ok": True}, "geometry": AREA},
            {"type": "Feature", "properties": {"name": "c"}, "geometry": POINT},
        ))
        self.assertEqual([r[0] for r in rows], ["7", "1", "2"])  # own id, else the position
        self.assertEqual(rows[0][1:], ({"name": "a", "n": 1}, AREA))
        self.assertEqual(stats.count, 3)
        self.assertEqual(stats.geometry_type(), "MultiPolygon")
        self.assertEqual(stats.field_types(), {"name": "null|string", "n": "number", "ok": "boolean"})

    def test_missing_geometry_and_properties(self):
        rows, stats = self.rows(collection({"type": "Feature", "properties": None, "geometry": None}))
        self.assertEqual(rows, [("0", {}, None)])
        self.assertIsNone(stats.geometry_type())

    def test_empty_collection(self):
        rows, stats = self.rows(collection())
        self.assertEqual((rows, stats.count), ([], 0))

    def test_not_a_feature_collection(self):
        # e.g. the forest-areas .json files: a plain list of attribute rows
        with self.assertRaises(ValueError):
            self.rows([{"id": 1}])
        with self.assertRaises(ValueError):
            self.rows(None)  # sync.load_json gives None for unreadable JSON

    def test_crs(self):
        wgs = {"type": "name", "properties": {"name": "urn:ogc:def:crs:OGC:1.3:CRS84"}}
        self.assertEqual(len(self.rows(collection({"geometry": POINT}, crs=wgs))[0]), 1)
        bgs = {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::7801"}}
        with self.assertRaises(ValueError):
            self.rows(collection({"geometry": POINT}, crs=bgs))


def rows_of(layers):
    return [(suffix, list(rows)) for suffix, rows in layers]


class TableReadersTest(unittest.TestCase):
    def test_csv_with_coordinates(self):
        # green-islands: trailing comma -> unnamed empty column
        data = "rayon,adres,latitude,longitude,\nОвча купел,бл. 1,42.669260,23.263460,\n,,,,\n".encode()
        [(suffix, rows)] = rows_of(load_db.parse_file("x.csv", data))
        self.assertEqual(suffix, "")
        self.assertEqual(rows, [("1", {"rayon": "Овча купел", "adres": "бл. 1", "latitude": "42.669260",
                                       "longitude": "23.263460"},
                                 {"type": "Point", "coordinates": [23.26346, 42.66926]})])

    def test_csv_cp1251_semicolon_decimal_comma(self):
        data = "адрес;lat;long\nул. Калоян 5;42,69543;23,322409\n".encode("cp1251")
        [(_, rows)] = rows_of(load_db.parse_file("x.csv", data))
        self.assertEqual(rows[0][1]["адрес"], "ул. Калоян 5")
        self.assertEqual(rows[0][2]["coordinates"], [23.322409, 42.69543])

    def test_no_or_bad_coordinates(self):
        [(_, rows)] = rows_of(load_db.parse_file("x.csv", b"a,b\n1,2\n"))
        self.assertIsNone(rows[0][2])
        [(_, rows)] = rows_of(load_db.parse_file("x.csv", b"lat,lon\n,\nn/a,23.3\n0,0\n"))
        self.assertEqual([r[2] for r in rows], [None, None])

    def test_coordinate_column_names(self):
        names = ["Район", "Latitude (географска ширина)", "Longtitude (географска дължина)", "lateral"]
        self.assertEqual(load_db.coord_columns(names), (names[1], names[2]))
        self.assertEqual(load_db.coord_columns(["lat", "long"]), ("lat", "long"))
        self.assertEqual(load_db.coord_columns(["latitude", "lng"]), ("latitude", "lng"))
        self.assertEqual(load_db.coord_columns(["platform", "along"]), (None, None))

    def test_table_rows_header_and_padding(self):
        # preferential-parking-disability-cards.xls: blank first column, dates as text
        rows = [[None, None], ["", "Номер", "Валидна до", "Номер"], ["", 133950.0, "01.05.2028", 1.0],
                ["", None, None, None], ["", 5.0, None, None, "extra"]]
        self.assertEqual(load_db.table_rows(rows), [
            {"Номер": 133950, "Валидна до": "01.05.2028", "Номер_2": 1},
            {"Номер": 5, "Валидна до": None, "Номер_2": None, "col_5": "extra"},
        ])
        self.assertEqual(load_db.table_rows([]), [])

    def test_cells(self):
        self.assertEqual(load_db.cell(dt.datetime(2026, 1, 2, 3, 4)), "2026-01-02T03:04:00")
        self.assertEqual(load_db.cell(1.5), 1.5)
        self.assertIs(load_db.cell(True), True)

    def test_json_list(self):
        data = json.dumps([{"id": 1, "n_type": "Бял бор"}, {"id": 2, "n_type": None}]).encode()
        [(_, rows)] = rows_of(load_db.parse_file("x.json", data))
        self.assertEqual(rows, [("1", {"id": 1, "n_type": "Бял бор"}, None), ("2", {"id": 2, "n_type": None}, None)])

    def test_json_feature_collection(self):
        data = json.dumps(collection({"id": "a", "properties": {}, "geometry": POINT})).encode("utf-8-sig")
        [(_, rows)] = rows_of(load_db.parse_file("x.geojson", data))
        self.assertEqual(rows, [("a", {}, POINT)])

    def workbook(self, *sheets):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, rows in sheets:
            ws = wb.create_sheet(title)
            for r in rows:
                ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_xlsx_single_sheet(self):
        data = self.workbook(("Sheet1", [["id", "lat", "long"], ["SOFARM-001", 42.689355, 23.318895]]))
        self.assertEqual(rows_of(load_db.parse_file("x.xlsx", data)), [
            ("", [("1", {"id": "SOFARM-001", "lat": 42.689355, "long": 23.318895},
                   {"type": "Point", "coordinates": [23.318895, 42.689355]})])])

    def test_xlsx_one_layer_per_non_empty_sheet(self):
        data = self.workbook(("A", [["x", "y"], [1, 2]]), ("Empty", []), ("B", [["x", "y"], [3, 4]]))
        layers = rows_of(load_db.parse_file("x.xlsx", data))
        self.assertEqual([s for s, _ in layers], ["!/A", "!/B"])
        self.assertEqual(layers[1][1], [("1", {"x": 3, "y": 4}, None)])


class KindTest(unittest.TestCase):
    def test_kinds(self):
        self.assertEqual(load_db.kind_of(["a/b.geojson"], "ok"), "vector")
        self.assertEqual(load_db.kind_of(["x.shp", "x.dbf", "x.prj"], "ok"), "vector")
        self.assertEqual(load_db.kind_of(["x.csv"], "ok"), "table")
        self.assertEqual(load_db.kind_of(["x.7z"], "ok"), "archive")
        self.assertEqual(load_db.kind_of(["https://www.sofia.bg/page"], "link"), "link")
        self.assertEqual(load_db.kind_of(["https://gtfs.sofiatraffic.bg/api/v1/alerts"], "ok"), "other")


class CatalogRowsTest(unittest.TestCase):
    pkg = {
        "id": "pkg-1", "name": "bus-lines", "title": "Bus lines",
        "groups": [{"name": "mobility"}, {"name": "buldings"}],
        "organization": {"name": "sofiaplan"},
        "tags": [{"name": "GeoJSON"}, {"name": "CSV"}],
        "extras": [{"key": "Актуален към", "value": "28.08.2026 г."}],
        "license_id": "cc-by",
        "metadata_created": "2026-01-01T00:00:00.5",
        "metadata_modified": "2026-02-01T10:00:00+00:00",
    }

    def test_dataset_row(self):
        row = load_db.dataset_row(self.pkg)
        self.assertEqual(row["section"], "buildings")  # misspelled group slug fixed, as in sync.py
        self.assertEqual(row["groups"], ["buildings", "mobility"])
        self.assertEqual(row["tags"], ["CSV", "GeoJSON"])
        self.assertEqual(row["organization"], "sofiaplan")
        self.assertEqual(row["license"], "cc-by")
        self.assertEqual(row["extras"], {"Актуален към": "28.08.2026 г."})
        self.assertEqual(row["metadata_created"], "2026-01-01T00:00:00.5+00:00")
        self.assertEqual(row["metadata_modified"], "2026-02-01T10:00:00+00:00")
        self.assertIs(row["raw"], self.pkg)

    def test_resource_row_multi_part(self):
        meta = {"status": "ok", "size": 3, "parts": [{"file": "r__a.geojson", "sha256": "aa"},
                                                     {"file": "r__b.geojson", "sha256": "bb"}]}
        res = {"id": "r", "url": "https://a/1, https://a/2", "format": "GeoJSON"}
        row = load_db.resource_row(self.pkg, res, meta)
        self.assertEqual(row["file_path"], "buildings/bus-lines/r__a.geojson;buildings/bus-lines/r__b.geojson")
        self.assertEqual(row["kind"], "vector")
        self.assertEqual(load_db.resource_layers(meta, "buildings/bus-lines"),
                         [("buildings/bus-lines/r__a.geojson", "aa"), ("buildings/bus-lines/r__b.geojson", "bb")])

    def test_resource_row_not_downloaded(self):
        res = {"id": "r", "url": "https://api.sofiaplan.bg/datasets/1", "format": "GeoJSON"}
        row = load_db.resource_row(self.pkg, res, {"status": "error", "file": "stale.geojson"})
        self.assertEqual((row["file_path"], row["kind"], row["status"]), (None, "vector", "error"))
        self.assertEqual(load_db.resource_layers({"status": "error"}, "x"), [])
        self.assertEqual(load_db.resource_row(self.pkg, res, {})["status"], "not-fetched")

    def test_only_readable_files_become_layers(self):
        self.assertEqual(load_db.resource_layers({"status": "ok", "file": "r__x.7z", "sha256": "cc"}, "d"), [])
        meta = {"status": "ok", "parts": [{"file": "r__x.CSV", "sha256": "a"}, {"file": "r__x.pdf", "sha256": "b"},
                                          {"file": "r__x.xlsx", "sha256": "c"}]}
        self.assertEqual(load_db.resource_layers(meta, "d"), [("d/r__x.CSV", "a"), ("d/r__x.xlsx", "c")])


if __name__ == "__main__":
    unittest.main()
